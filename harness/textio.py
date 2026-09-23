"""Извлечение текста из файлов (для инструментов и RAG)."""
from pathlib import Path

TEXT_SUFFIXES = {".txt", ".md", ".csv", ".json", ".html", ".htm", ".xml", ".log", ".yaml", ".yml", ".tsv"}
DOC_SUFFIXES = TEXT_SUFFIXES | {".pdf", ".docx", ".xlsx", ".xls"}


def extract_text(path: Path, max_rows=300) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader
        return "\n\n".join(f"[стр. {i + 1}]\n{p.extract_text() or ''}"
                           for i, p in enumerate(PdfReader(path).pages))
    if suffix == ".docx":
        from docx import Document
        doc = Document(path)
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    if suffix in (".xlsx", ".xls"):
        import pandas as pd
        sheets = pd.read_excel(path, sheet_name=None)
        return "\n\n".join(f"[лист {name}]\n{df.head(max_rows).to_csv(index=False)}"
                           for name, df in sheets.items())
    if suffix in (".html", ".htm"):
        from bs4 import BeautifulSoup
        return BeautifulSoup(path.read_text(encoding="utf-8", errors="replace"), "html.parser").get_text("\n")
    return path.read_text(encoding="utf-8", errors="replace")


def chunk_text(text, size=1200, overlap=200):
    """Режем по абзацам на куски ~size символов с перекрытием."""
    paragraphs = [p.strip() for p in text.split("\n") if p.strip()]
    chunks, current = [], ""
    for p in paragraphs:
        while len(p) > size:  # очень длинный абзац
            if current:
                chunks.append(current)
                current = ""
            chunks.append(p[:size])
            p = p[size - overlap:]
        if len(current) + len(p) + 1 > size and current:
            chunks.append(current)
            current = current[-overlap:] + "\n" + p
        else:
            current = f"{current}\n{p}" if current else p
    if current.strip():
        chunks.append(current)
    return chunks
